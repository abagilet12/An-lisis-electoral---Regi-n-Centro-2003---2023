#!/usr/bin/env python3
"""
Genera un .xlsx por eleccion con TODO lo recolectado para esa eleccion.

A diferencia de la version anterior, que solo emitia el nivel mas fino, este
incluye **todos los niveles disponibles** en hojas separadas, mas la
clasificacion politica homologada, la procedencia de cada dato y los
criterios de construccion, de modo que cada archivo se entienda solo.

Hojas:
  Criterios        las reglas de construccion, para no depender de otro archivo
  Resumen          totales de la eleccion
  Fuentes          que fuente alimento cada nivel y con que tipo de recuento
  Circuito_ancho   una fila por circuito, una columna por agrupacion (mapas)
  Circuito_largo   una fila por circuito y agrupacion, con familia y bloque
  Departamento     agregado por departamento, si hay
  Localidad        agregado por localidad, si hay
  Provincia        total provincial, con escrutinio definitivo si hay
  Nomenclador      circuitos, con padron, mesas y si tienen poligono

Salida: salida/xlsx/<anio>_<instancia>_santa_fe_presidente.xlsx
"""
import csv
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
PROC = RAIZ / "datos" / "procesados"
GEO = RAIZ / "datos" / "geo" / "circuitos"
SALIDA = RAIZ / "salida" / "xlsx"
TABLA_HOM = RAIZ / "datos" / "referencia" / "homologacion_agrupaciones.csv"

FUENTE = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
NORMAL = Font(name=FUENTE, size=10)
TITULO = Font(name=FUENTE, bold=True, size=11)

CARGO, DISTRITO_ID, DISTRITO = "PRESIDENTE Y VICE", "21", "Santa Fe"
NO_POSITIVOS = ["blancos", "nulos", "impugnados", "recurridos", "comando"]
ORDEN = {"PASO": 0, "GENERAL": 1, "BALOTAJE": 2}

FECHAS = {
    ("2003", "GENERAL"): "2003-04-27", ("2007", "GENERAL"): "2007-10-28",
    ("2011", "PASO"): "2011-08-14", ("2011", "GENERAL"): "2011-10-23",
    ("2015", "PASO"): "2015-08-09", ("2015", "GENERAL"): "2015-10-25",
    ("2015", "BALOTAJE"): "2015-11-22",
    ("2019", "PASO"): "2019-08-11", ("2019", "GENERAL"): "2019-10-27",
    ("2023", "PASO"): "2023-08-13", ("2023", "GENERAL"): "2023-10-22",
    ("2023", "BALOTAJE"): "2023-11-19",
}

# archivo -> (nivel, etiqueta legible de la fuente)
ORIGENES = [
    ("resultados_circuito_polar.csv", "circuito",
     "PoliticaArgentina/data_warehouse (Atlas Electoral de Andy Tow)"),
    ("2023_PASO_circuito.csv", "circuito", "DINE, archivo por mesa"),
    ("2023_GENERAL_circuito.csv", "circuito", "DINE, archivo por mesa"),
    ("2023_BALOTAJE_circuito.csv", "circuito", "DINE, archivo por mesa"),
    ("resultados_departamento.csv", "departamento", "API de la DINE"),
    ("serie_localidad.csv", "localidad",
     "Archivos «Votos por Localidad» y agregado desde mesa"),
    ("resultados.csv", "provincia",
     "Consulta de Escrutinio por Zona (Justicia Nacional Electoral)"),
    ("resultados_provincia_dine.csv", "provincia",
     "Archivos nacionales de la DINE"),
    ("resultados_provincia_transcripto.csv", "provincia",
     "Atlas Electoral de Andy Tow, transcripto"),
]

CRITERIOS = [
    ("Alcance", "Categoría Presidente y Vice, provincia de Santa Fe "
     "(distrito 21), 2003-2023. Doce instancias de votación."),
    ("Unidad de referencia", "El circuito electoral. Es la menor unidad con "
     "identidad geográfica estable y es la que se usa para mapear. "
     "Departamento y provincia se obtienen sumando circuitos."),
    ("Al sumar votos", "Filtrar SIEMPRE tipo_registro = 'agrupacion'. Sumar "
     "sin filtrar cuenta dos veces: las listas internas de las PASO ya están "
     "contenidas en su agrupación, y 'positivo' y 'total' son agregados."),
    ("Identificadores", "circuito_id y distrito_id son TEXTO, con los ceros "
     "a la izquierda de la fuente. '00115' no es '115'. Al leerlos con código "
     "hay que forzar el tipo o el cruce con el mapa falla en silencio."),
    ("Códigos entre años", "Los códigos de circuito NO son estables entre "
     "años: Santa Fe renumeró en 2023. Nunca cruzar un año con el "
     "nomenclador de otro."),
    ("Procedencia", "Cada fila declara su fuente y su recuento_tipo. Nunca "
     "mezclar recuento provisorio con definitivo en un mismo cálculo: "
     "difieren en torno al 2 % y el signo no es constante."),
    ("Homologación", "familia es el linaje partidario; bloque es el eje "
     "kirchnerismo / no kirchnerismo. Son decisiones de investigación, "
     "editables, no hechos del dato."),
    ("Comparabilidad", "Una serie temporal exige etiquetas homologadas, "
     "continuidad política declarada Y el mismo universo geográfico. "
     "Comparar la provincia de un año con 30 localidades de otro mide un "
     "cambio de recorte, no un cambio político."),
    ("Porcentajes", "Guardados como fracción, con formato de celda de "
     "porcentaje. 0,25 se muestra como 25 %."),
    ("Valores derivados", "Son números calculados por el script, no fórmulas: "
     "una fórmula sin valor cacheado se lee vacía desde código. La "
     "reproducibilidad la da el script."),
    ("Controles aplicados", "Las agrupaciones suman el total de positivos; "
     "positivos + blancos + nulos + impugnados + recurridos da el total de "
     "votantes; la participación calculada coincide con la declarada; en PASO "
     "las listas internas suman el total de su agrupación."),
]


def leer(nombre):
    ruta = PROC / nombre
    return list(csv.DictReader(open(ruta, encoding="utf-8"))) if ruta.exists() else []


def entero(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def cargar_homologacion():
    if not TABLA_HOM.exists():
        return {}
    return {(r["anio"], r["instancia"], r["agrupacion_original"].strip()):
            (r["agrupacion_homologada"], r["familia"], r["bloque"])
            for r in csv.DictReader(open(TABLA_HOM, encoding="utf-8"))}


def cargar_geo(anio):
    ruta = GEO / f"santafe_circuitos_{anio}_reconstruido.geojson"
    if not ruta.exists():
        return {}
    d = json.loads(ruta.read_text(encoding="utf-8"))
    return {f["properties"]["circuito"]: f["properties"].get("origen", "")
            for f in d["features"]}


# --------------------------------------------------------------------------
# Escritura de hojas
# --------------------------------------------------------------------------

def encabezar(ws, columnas, anchos=None):
    ws.append(columnas)
    for i, _ in enumerate(columnas, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (
            anchos[i - 1] if anchos else 14)
    ws.freeze_panes = "A2"


def formatear(ws, desde=2):
    for fila in ws.iter_rows(min_row=desde):
        for c in fila:
            c.font = NORMAL


def hoja_clave_valor(wb, titulo, pares, ancho=(34, 104)):
    ws = wb.create_sheet(titulo)
    encabezar(ws, ["concepto", "detalle"], list(ancho))
    for k, v in pares:
        ws.append([k, str(v)])
    for fila in ws.iter_rows(min_row=2):
        fila[0].font = Font(name=FUENTE, bold=True, size=10)
        fila[1].font = NORMAL
        fila[1].alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def hoja_larga(wb, titulo, filas, columnas, anchos):
    ws = wb.create_sheet(titulo)
    encabezar(ws, columnas, anchos)
    for f in filas:
        ws.append([f.get(c, "") for c in columnas])
    formatear(ws)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if c.column_letter and columnas[c.column - 1] in ("votos", "electores",
                                                              "mesas", "votantes"):
                c.number_format = "#,##0"
    return ws


def construir(anio, instancia, datos, hom, geo):
    wb = Workbook()
    wb.remove(wb.active)
    fecha = FECHAS.get((anio, instancia), "")

    def clasificar(etiqueta):
        h = hom.get((anio, instancia, etiqueta.strip()))
        return h if h else ("", "", "")

    # ---------------- Criterios ----------------
    hoja_clave_valor(wb, "Criterios", CRITERIOS)

    # ---------------- Circuito ----------------
    circ = datos.get("circuito", [])
    agrupaciones, votos_c, unidades_c = [], defaultdict(int), {}
    if circ:
        tot = defaultdict(int)
        for r in circ:
            if r["tipo_registro"] == "agrupacion":
                tot[r["agrupacion"].strip()] += entero(r["votos"])
        agrupaciones = [a for a, _ in sorted(tot.items(), key=lambda kv: -kv[1])]
        for r in circ:
            cid = r["circuito_id"]
            unidades_c.setdefault(cid, {
                "seccion_id": r.get("seccion_id", ""),
                "seccion": r.get("seccion", ""),
                "localidad": r.get("localidad", ""),
                "electores": entero(r.get("electores", 0)),
                "mesas": entero(r.get("mesas", 0))})
            if r["tipo_registro"] == "agrupacion":
                votos_c[(cid, r["agrupacion"].strip())] += entero(r["votos"])
            elif r["tipo_registro"] in NO_POSITIVOS:
                votos_c[(cid, r["tipo_registro"])] += entero(r["votos"])

        ws = wb.create_sheet("Circuito_ancho")
        cols = (["anio", "instancia", "fecha", "cargo", "distrito_id", "distrito",
                 "seccion_id", "seccion", "circuito_id", "electores", "mesas"]
                + [a.lower().replace(" ", "_")[:55] for a in agrupaciones]
                + NO_POSITIVOS + ["positivos", "total_votantes", "participacion",
                                  "tiene_poligono"])
        encabezar(ws, cols, [8, 10, 11, 18, 10, 12, 10, 18, 12, 12, 8]
                  + [16] * len(agrupaciones) + [11] * 9)
        for cid, u in sorted(unidades_c.items(),
                             key=lambda kv: (kv[1]["seccion"], kv[0])):
            va = [votos_c.get((cid, a), 0) for a in agrupaciones]
            vn = [votos_c.get((cid, t), 0) for t in NO_POSITIVOS]
            pos, total = sum(va), sum(va) + sum(vn)
            ws.append([anio, instancia, fecha, CARGO, DISTRITO_ID, DISTRITO,
                       u["seccion_id"], u["seccion"], cid,
                       u["electores"] or None, u["mesas"] or None]
                      + va + vn + [pos, total,
                                   (total / u["electores"]) if u["electores"] else None,
                                   "sí" if cid in geo else "no"])
        formatear(ws)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                if c.column >= 10 and c.column < len(cols):
                    c.number_format = "#,##0"
            fila[len(cols) - 2].number_format = "0.00%"
            for i in (4, 6, 8):
                fila[i].number_format = "@"

        largo = []
        for cid, u in sorted(unidades_c.items(),
                             key=lambda kv: (kv[1]["seccion"], kv[0])):
            for et in agrupaciones + NO_POSITIVOS:
                homol, fam, blo = (clasificar(et) if et in agrupaciones
                                   else ("", "", ""))
                largo.append({
                    "anio": anio, "instancia": instancia, "fecha": fecha,
                    "cargo": CARGO, "distrito_id": DISTRITO_ID,
                    "distrito": DISTRITO, "seccion_id": u["seccion_id"],
                    "seccion": u["seccion"], "circuito_id": cid,
                    "agrupacion": et,
                    "tipo_registro": "agrupacion" if et in agrupaciones else et,
                    "votos": votos_c.get((cid, et), 0),
                    "agrupacion_homologada": homol, "familia": fam,
                    "bloque": blo})
        hoja_larga(wb, "Circuito_largo", largo,
                   ["anio", "instancia", "fecha", "cargo", "distrito_id",
                    "distrito", "seccion_id", "seccion", "circuito_id",
                    "agrupacion", "tipo_registro", "votos",
                    "agrupacion_homologada", "familia", "bloque"],
                   [8, 10, 11, 18, 10, 12, 10, 18, 12, 42, 14, 12, 40, 30, 26])

    # ---------------- Departamento / Localidad / Provincia ----------------
    def bloque_nivel(nombre, filas, claves, anchos):
        if not filas:
            return
        out = []
        for r in filas:
            et = r["agrupacion"].strip()
            homol, fam, blo = (clasificar(et)
                               if r["tipo_registro"] == "agrupacion"
                               else ("", "", ""))
            out.append({**{k: r.get(k, "") for k in claves},
                        "agrupacion_homologada": homol, "familia": fam,
                        "bloque": blo})
        hoja_larga(wb, nombre, out, claves + ["agrupacion_homologada",
                                              "familia", "bloque"],
                   anchos + [40, 30, 26])

    bloque_nivel("Departamento", datos.get("departamento", []),
                 ["anio", "instancia", "seccion_id", "mesas", "electores",
                  "votantes", "agrupacion", "tipo_registro", "votos",
                  "recuento_tipo"],
                 [8, 10, 11, 9, 12, 12, 42, 14, 12, 15])
    bloque_nivel("Localidad", datos.get("localidad", []),
                 ["anio", "instancia", "seccion", "localidad", "circuito_id",
                  "agrupacion", "tipo_registro", "votos", "recuento_tipo"],
                 [8, 10, 20, 26, 12, 42, 14, 12, 15])
    bloque_nivel("Provincia", datos.get("provincia", []),
                 ["anio", "instancia", "agrupacion", "formula",
                  "tipo_registro", "votos", "recuento_tipo"],
                 [8, 10, 46, 40, 14, 12, 15])

    # ---------------- Nomenclador ----------------
    if unidades_c:
        ws = wb.create_sheet("Nomenclador")
        encabezar(ws, ["circuito_id", "seccion_id", "seccion", "localidad",
                       "mesas", "electores", "tiene_poligono", "origen_poligono"],
                  [12, 11, 20, 26, 9, 12, 14, 18])
        for cid, u in sorted(unidades_c.items(),
                             key=lambda kv: (kv[1]["seccion"], kv[0])):
            ws.append([cid, u["seccion_id"], u["seccion"], u["localidad"],
                       u["mesas"] or None, u["electores"] or None,
                       "sí" if cid in geo else "no", geo.get(cid, "")])
        formatear(ws)
        for fila in ws.iter_rows(min_row=2):
            fila[0].number_format = fila[1].number_format = "@"
            fila[4].number_format = fila[5].number_format = "#,##0"

    # ---------------- Fuentes ----------------
    pares = []
    for niv in ["circuito", "departamento", "localidad", "provincia"]:
        fs = datos.get(niv, [])
        if not fs:
            pares.append((f"Nivel {niv}", "No disponible para esta elección."))
            continue
        etiqueta = datos["etiquetas"].get(niv, "")
        rec = {f.get("recuento_tipo", "") for f in fs} or {""}
        unidades = len({f.get("circuito_id") or f.get("seccion_id")
                        or f.get("localidad") or "provincia" for f in fs})
        pares.append((f"Nivel {niv}",
                      f"{etiqueta}. {unidades} unidades. "
                      f"Recuento: {', '.join(sorted(x for x in rec if x)) or 'no declarado'}."))
    if circ:
        con = sum(1 for cid in unidades_c if cid in geo)
        pares.append(("Cartografía",
                      f"{con} de {len(unidades_c)} circuitos con polígono. "
                      "Capas de Franco Galeano, licencia CC BY: citar en "
                      "cualquier mapa publicado."))
    hoja_clave_valor(wb, "Fuentes", pares)

    # ---------------- Resumen ----------------
    ws = wb.create_sheet("Resumen", 0)
    ws.append(["Elección", f"{instancia} {anio}"])
    ws.append(["Fecha", fecha])
    ws.append(["Cargo", CARGO])
    ws.append(["Distrito", f"{DISTRITO} (id {DISTRITO_ID})"])
    ws.append(["Niveles disponibles",
               ", ".join(n for n in ["circuito", "departamento", "localidad",
                                     "provincia"] if datos.get(n))])
    ws.append([])
    fila_enc = ws.max_row + 1
    ws.append(["agrupación", "familia", "bloque", "votos", "% positivos"])
    if circ:
        tot = {a: sum(votos_c.get((c, a), 0) for c in unidades_c)
               for a in agrupaciones}
        nop = {t: sum(votos_c.get((c, t), 0) for c in unidades_c)
               for t in NO_POSITIVOS}
        pos = sum(tot.values())
        for a in agrupaciones:
            _, fam, blo = clasificar(a)
            ws.append([a, fam, blo, tot[a], tot[a] / pos if pos else None])
        for t in NO_POSITIVOS:
            ws.append([t, "", "", nop[t], None])
        ws.append(["VOTOS POSITIVOS", "", "", pos, None])
        ws.append(["TOTAL VOTANTES", "", "", pos + sum(nop.values()), None])
        padron = sum(u["electores"] for u in unidades_c.values())
        ws.append(["ELECTORES INSCRIPTOS", "", "", padron, None])
        ws.append(["PARTICIPACIÓN", "", "",
                   (pos + sum(nop.values())) / padron if padron else None, None])
    formatear(ws, 1)
    for i in (1, 2, 3, 4, 5, fila_enc):
        for c in ws[i]:
            c.font = TITULO if i < 6 else HEAD_FONT
            if i == fila_enc:
                c.fill = HEAD_FILL
    for fila in ws.iter_rows(min_row=fila_enc + 1):
        fila[3].number_format = "#,##0"
        fila[4].number_format = "0.00%"
    if ws.max_row > fila_enc:
        ws.cell(row=ws.max_row, column=4).number_format = "0.00%"
    for col, a in (("A", 40), ("B", 32), ("C", 26), ("D", 16), ("E", 13)):
        ws.column_dimensions[col].width = a

    SALIDA.mkdir(parents=True, exist_ok=True)
    destino = SALIDA / f"{anio}_{instancia}_santa_fe_presidente.xlsx"
    wb.save(destino)
    return destino, [s for s in wb.sheetnames]


def main():
    hom = cargar_homologacion()
    por_eleccion = defaultdict(lambda: defaultdict(list))
    etiquetas = defaultdict(dict)
    for archivo, nivel, etiqueta in ORIGENES:
        for r in leer(archivo):
            k = (r["anio"], r["instancia"])
            if por_eleccion[k][nivel] and etiquetas[k].get(nivel) != etiqueta:
                continue          # un nivel se toma de una sola fuente
            por_eleccion[k][nivel].append(r)
            etiquetas[k][nivel] = etiqueta

    print(f"{'elección':17}{'circ.':>7}{'hojas':>7}  archivo")
    for k in sorted(por_eleccion, key=lambda x: (x[0], ORDEN[x[1]])):
        datos = dict(por_eleccion[k])
        datos["etiquetas"] = etiquetas[k]
        destino, hojas = construir(k[0], k[1], datos, hom, cargar_geo(k[0]))
        nc = len({r["circuito_id"] for r in datos.get("circuito", [])})
        print(f"  {k[0]} {k[1]:10}{nc:>7}{len(hojas):>7}  {destino.name}")
    print(f"\n{len(por_eleccion)} archivos en salida/xlsx/")


if __name__ == "__main__":
    main()
