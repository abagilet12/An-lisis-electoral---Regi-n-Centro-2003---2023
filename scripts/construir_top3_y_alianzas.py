#!/usr/bin/env python3
"""
Top 3 por elección y composición de las alianzas, 2003-2023.

Genera dos CSV en datos/referencia/:

  top3_por_eleccion.csv    los tres primeros por instancia, con la etiqueta
                           tal cual la publica cada fuente.
  composicion_alianzas.csv partidos integrantes de cada alianza, copiados del
                           encabezado de los Excel oficiales de la DINE
                           (disponible 2003, 2007 y 2011).

Fuentes (todas del Ministerio del Interior / DINE):
  - Excel 2003-2015 y 2019 de argentina.gob.ar/dine/resultados-electorales
  - CSV de mesas de las generales 2023 (zip en la misma página)
  - API resultados.mininterior.gob.ar (PASO y balotaje 2023, provisorio)
  - Base propia: datos/procesados/serie_homologada.csv (recuento provisorio)

Uso: python3 scripts/construir_top3_y_alianzas.py [carpeta_cache]
Requiere openpyxl. Descarga lo que falte en la carpeta de caché.
"""
import csv
import collections
import json
import re
import subprocess
import sys
import unicodedata
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
CACHE = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "datos" / "crudos" / "dine_nacional"
SALIDA = RAIZ / "datos" / "referencia"
WEB = "https://www.argentina.gob.ar/sites/default/files/"
API = "https://resultados.mininterior.gob.ar/api/resultados/getResultados"

# (anio, instancia, archivo, hoja nacional, hoja Santa Fe, formato)
EXCEL = [
    (2003, "GENERAL", "resultados_2003_presidente_y_vicepresidente.xlsx", "Nacionales", "Santa Fe", "A"),
    (2007, "GENERAL", "resultados_2007_presidente_y_vicepresidente.xlsx", "Nacionales", "Santa Fe", "A"),
    (2011, "PASO", "resultados_2011_eleccion_paso_presidente_y_vicepresidente.xlsx", "Nacionales", "Santa Fe", "A"),
    (2011, "GENERAL", "resultados_2011_generales_presidente_y_vicepresidente.xlsx", "Nacionales", "Santa Fe", "A"),
    (2015, "PASO", "resultados_2015_elecciones_paso_presidente_y_vicepresidente.xlsx", "Nacional", "Santa Fe", "B"),
    (2015, "GENERAL", "resultados_2015_nacionales_presidente_y_vicepresidente.xlsx", "Nacional", "Santa Fe", "B"),
    (2015, "BALOTAJE", "resultados_2015_segunda_vuelta_presidente_y_vicepresidente.xlsx", "Nacional", "Santa Fe", "B"),
]


def bajar(nombre):
    destino = CACHE / nombre
    if not destino.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(WEB + nombre, destino)
    return destino


def norm(s):
    s = unicodedata.normalize("NFD", str(s)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z ]", "", s).strip()


def filas(ws):
    return [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)
            if any(c is not None for c in r)]


def leer_resultados(hoja, formato):
    """Devuelve (lista [(formula, etiqueta, votos)], positivos)."""
    out, positivos = [], None
    for f in filas(hoja):
        if isinstance(f[0], str) and f[0].strip().upper().startswith("VOTOS POSITIVOS"):
            positivos = f[1]
            break
        if formato == "B":                      # [codigo, etiqueta, votos, pct]
            if isinstance(f[0], int) and len(f) >= 3 and isinstance(f[2], (int, float)):
                out.append(("", str(f[1]).strip(), int(f[2])))
        else:                                   # A: [formula, (agrupacion,) votos, %]
            nums = [c for c in f if isinstance(c, (int, float))]
            textos = [c for c in f if isinstance(c, str)]
            if len(nums) == 2 and nums[0] > 100 and textos and " - " in textos[0]:
                etiqueta = textos[1].strip() if len(textos) > 1 else ""
                out.append((textos[0].strip(), etiqueta, int(nums[0])))
    return out, positivos


def etiquetas_por_formula(hoja):
    """Hoja nacional 2003/2007: tabla formula -> agrupacion del encabezado."""
    m = {}
    for f in filas(hoja):
        if isinstance(f[0], str) and f[0].strip().upper().startswith("TOTAL PA"):
            break
        if len(f) == 2 and " - " in f[0] and f[0].upper() == f[0]:
            m[norm(f[0].split(",")[0])] = f[1].strip()
    return m


def composicion(hoja, anio, instancia):
    """Alianzas y partidos integrantes del encabezado de la hoja nacional."""
    out, actual = [], None
    for f in filas(hoja):
        t = f[0] if isinstance(f[0], str) else None
        if t is None:
            continue
        if t.strip().upper().startswith("TOTAL PA"):
            break
        if t.strip().upper().startswith("ALIANZA ") and t.strip() == t.strip().upper() and len(f) == 1:
            actual = t.strip()
        elif actual and len(f) == 1 and not t.isupper():
            out.append((anio, instancia, actual, t.strip()))
    return out


def top3_oficial(ambito, anio, inst, res, positivos, fuente, recuento, etq_formula=None):
    res = sorted(res, key=lambda x: -x[2])[:3]
    rows = []
    for i, (formula, etiqueta, votos) in enumerate(res, 1):
        if not etiqueta and etq_formula:
            etiqueta = etq_formula.get(norm(formula.split("-")[0]), "")
        rows.append([ambito, fuente, recuento, anio, inst, i, etiqueta, formula,
                     votos, round(100 * votos / positivos, 2)])
    return rows


def main():
    top3, comp = [], []
    # --- Excel DINE 2003-2015 (escrutinio definitivo) ---
    for anio, inst, nombre, hn, hs, fmt in EXCEL:
        wb = openpyxl.load_workbook(bajar(nombre), read_only=True, data_only=True)
        etq = etiquetas_por_formula(wb[hn]) if anio in (2003, 2007) else None
        for ambito, hoja in (("Nacional", hn), ("Santa Fe", hs)):
            res, pos = leer_resultados(wb[hoja], fmt)
            top3 += top3_oficial(ambito, anio, inst, res, pos, "DINE (Excel oficial)", "DEFINITIVO", etq)
        if anio <= 2011:
            comp += composicion(wb[hn], anio, inst)

    # --- 2019: Excel definitivo (total país PASO y generales; Santa Fe generales) ---
    wb = openpyxl.load_workbook(bajar("2019_pv_definitivos_total_pais_paso_y_generales_1.xlsx"), read_only=True, data_only=True)
    for inst, hoja in (("PASO", "TOTAL PAIS PASO"), ("GENERAL", "TOTAL PAIS GENERALES")):
        res, pos = leer_resultados(wb[hoja], "B")
        top3 += top3_oficial("Nacional", 2019, inst, res, pos, "DINE (Excel oficial)", "DEFINITIVO")
    wb = openpyxl.load_workbook(bajar("p.v._definitivo_x_distrito_grales_2019.xlsx"), read_only=True, data_only=True)
    res, pos = leer_resultados(wb["SANTA FE"], "B")
    top3 += top3_oficial("Santa Fe", 2019, "GENERAL", res, pos, "DINE (Excel oficial)", "DEFINITIVO")

    # --- 2023: API (PASO y balotaje, país) y CSV de mesas (generales) ---
    for inst, tipo in (("PASO", 1), ("BALOTAJE", 3)):
        url = f"{API}?anioEleccion=2023&tipoRecuento=1&tipoEleccion={tipo}&categoriaId=1"
        d = json.load(urllib.request.urlopen(url, timeout=60))
        v = [(x["nombreAgrupacion"], x["votos"]) for x in d["valoresTotalizadosPositivos"]]
        pos = sum(x[1] for x in v)
        for i, (n, vt) in enumerate(sorted(v, key=lambda x: -x[1])[:3], 1):
            top3.append(["Nacional", "DINE (API)", "PROVISORIO", 2023, inst, i, n, "", vt, round(100 * vt / pos, 2)])
    z = bajar("2023_generales_1.zip")
    nac, sf = collections.Counter(), collections.Counter()
    with zipfile.ZipFile(z) as zf, zf.open("2023_Generales/ResultadoElectorales_2023_Generales.csv") as fh:
        import io
        for x in csv.DictReader(io.TextIOWrapper(fh, encoding="utf8")):
            if x["votos_tipo"] == "POSITIVO" and x["cargo_id"] == "1":
                n = int(x["votos_cantidad"])
                nac[x["agrupacion_nombre"]] += n
                if x["distrito_id"] == "21":
                    sf[x["agrupacion_nombre"]] += n
    for ambito, c in (("Nacional", nac), ("Santa Fe", sf)):
        pos = sum(c.values())
        for i, (n, vt) in enumerate(c.most_common(3), 1):
            top3.append([ambito, "DINE (CSV de mesas)", "PROVISORIO", 2023, "GENERAL", i, n, "", vt, round(100 * vt / pos, 2)])

    # --- Base propia (Santa Fe, recuento provisorio, las 12 instancias) ---
    tot = collections.defaultdict(lambda: collections.defaultdict(int))
    fam = {}
    with open(RAIZ / "datos" / "procesados" / "serie_homologada.csv", encoding="utf8") as fh:
        for r in csv.DictReader(fh):
            if r["nivel"] == "circuito":
                k = (int(r["anio"]), r["instancia"])
                tot[k][r["agrupacion_original"]] += int(float(r["votos"]))
                fam[(k, r["agrupacion_original"])] = r["agrupacion_homologada"]
    for k, d in tot.items():
        pos = sum(d.values())
        for i, (n, vt) in enumerate(sorted(d.items(), key=lambda x: -x[1])[:3], 1):
            top3.append(["Santa Fe", "Base propia (PolAr / DINE)", "PROVISORIO", k[0], k[1], i, n,
                         fam[(k, n)], vt, round(100 * vt / pos, 2)])

    orden = {"GENERAL": 1, "PASO": 0, "BALOTAJE": 2}
    top3.sort(key=lambda r: (r[0] != "Nacional", r[3], orden[r[4]], r[1], r[5]))
    SALIDA.mkdir(parents=True, exist_ok=True)
    with open(SALIDA / "top3_por_eleccion.csv", "w", newline="", encoding="utf8") as fh:
        w = csv.writer(fh)
        w.writerow(["ambito", "fuente", "recuento", "anio", "instancia", "puesto", "etiqueta_tal_cual",
                    "formula_o_homologada", "votos", "pct_votos_positivos"])
        w.writerows(top3)
    with open(SALIDA / "composicion_alianzas.csv", "w", newline="", encoding="utf8") as fh:
        w = csv.writer(fh)
        w.writerow(["anio", "instancia", "alianza", "partido_integrante"])
        w.writerows(comp)
    print(len(top3), "filas top3;", len(comp), "filas de composición")


if __name__ == "__main__":
    main()
