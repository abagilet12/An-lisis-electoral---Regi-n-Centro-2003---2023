#!/usr/bin/env python3
"""
Normaliza los .xlsx "resultados_<anio>_<instancia>_presidente_y_vicepresidente"
de la Direccion Nacional Electoral.

Son archivos nacionales con una hoja por distrito; se lee la de Santa Fe. El
dato es provincial: no traen desagregacion por departamento ni circuito.

Su valor es doble: completan el nivel provincial en anios que la Consulta de
Escrutinio por Zona no cubre, y son la unica fuente que vincula la FORMULA
(nombres de los candidatos) con la AGRUPACION que la lleva, que es el puente
necesario para homologar la serie.

El formato cambia entre archivos, asi que el parser reconoce la fila por su
forma en vez de asumir columnas fijas:
  2011 PASO     [formula, agrupacion, votos, pct]
  2015 PASO     [id, agrupacion, votos, pct_pos, pct_val]
                [precandidatos, votos, pct_interna]   <- lista interna
  2015 BALOTAJE [id, agrupacion, votos, pct]

Salida:
  datos/procesados/resultados_provincia_dine.csv
  datos/procesados/formulas.csv   (el puente formula -> agrupacion)
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
CRUDOS = RAIZ / "datos" / "crudos" / "provincia_dine"
SALIDA = RAIZ / "datos" / "procesados"

HOJA = "Santa Fe"
DISTRITO_ID, DISTRITO, CARGO = "21", "Santa Fe", "PRESIDENTE Y VICE"

# Etiquetas del bloque de cierre -> tipo_registro.
CIERRE = {
    "VOTOS VALIDOS": "validos", "VOTOS POSITIVOS": "positivo",
    "VOTOS EN BLANCO": "blancos", "VOTOS NULOS": "nulos",
    "VOTOS RECURRIDOS": "recurridos", "VOTOS IMPUGNADOS": "impugnados",
    "TOTAL DE VOTANTES": "total", "ELECTORES INSCRIPTOS": "inscriptos",
    "PORCENTAJE DE VOTANTES": "participacion",
}

COLUMNAS = ["anio", "instancia", "cargo", "distrito_id", "distrito",
            "agrupacion_id", "agrupacion", "formula", "tipo_registro",
            "votos", "pct", "archivo_origen"]
COLUMNAS_FORMULAS = ["anio", "instancia", "agrupacion_id", "agrupacion",
                     "formula", "votos"]


def limpiar(fila):
    return [c for c in fila if c is not None and str(c).strip() != ""]


def normalizar(s):
    """Mayusculas sin tildes, para comparar etiquetas."""
    return (str(s).upper().strip()
            .replace("Á", "A").replace("É", "E").replace("Í", "I")
            .replace("Ó", "O").replace("Ú", "U"))


def es_codigo(v):
    return bool(re.fullmatch(r"\d{1,4}", str(v).strip().replace(".0", "")))


def numero(v):
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return ""


def parsear(archivo):
    m = re.match(r"(\d{4})_([A-Z]+)_", archivo.name)
    anio, instancia = m.group(1), m.group(2)
    comun = {"anio": anio, "instancia": instancia, "cargo": CARGO,
             "distrito_id": DISTRITO_ID, "distrito": DISTRITO,
             "archivo_origen": archivo.name}

    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"{archivo.name}: no tiene hoja '{HOJA}'")
    filas = [limpiar(f) for f in wb[HOJA].iter_rows(values_only=True)]
    wb.close()

    resultados, formulas = [], []
    agrupacion_actual = {"id": "", "nombre": ""}

    for fila in filas:
        if not fila:
            continue
        etiqueta = normalizar(fila[0])

        # Las filas de encabezado tienen la misma forma que las de datos: lo
        # que las distingue es que no traen un numero de votos.
        if etiqueta not in CIERRE and not any(
            isinstance(c, (int, float)) for c in fila[1:]
        ):
            continue

        if etiqueta in CIERRE and len(fila) >= 2:
            resultados.append({**comun, "agrupacion_id": "",
                               "agrupacion": str(fila[0]).strip(), "formula": "",
                               "tipo_registro": CIERRE[etiqueta],
                               "votos": numero(fila[1]),
                               "pct": numero(fila[2]) if len(fila) > 2 else ""})
            continue

        # Fila de agrupacion: abre con un codigo numerico (2015) o trae la
        # formula y la agrupacion en columnas separadas (2011).
        if es_codigo(fila[0]) and len(fila) >= 3:
            agrupacion_actual = {"id": str(fila[0]).strip().replace(".0", ""),
                                 "nombre": str(fila[1]).strip()}
            resultados.append({**comun, "agrupacion_id": agrupacion_actual["id"],
                               "agrupacion": agrupacion_actual["nombre"],
                               "formula": "", "tipo_registro": "agrupacion",
                               "votos": numero(fila[2]),
                               "pct": numero(fila[3]) if len(fila) > 3 else ""})
            continue

        if len(fila) >= 4 and not es_codigo(fila[0]):
            # 2011: la formula y la agrupacion vienen juntas en la misma fila.
            formula, agrupacion = str(fila[0]).strip(), str(fila[1]).strip()
            resultados.append({**comun, "agrupacion_id": "",
                               "agrupacion": agrupacion, "formula": formula,
                               "tipo_registro": "agrupacion",
                               "votos": numero(fila[2]), "pct": numero(fila[3])})
            formulas.append({"anio": anio, "instancia": instancia,
                             "agrupacion_id": "", "agrupacion": agrupacion,
                             "formula": formula, "votos": numero(fila[2])})
            continue

        if len(fila) == 3 and agrupacion_actual["nombre"]:
            # 2015 PASO: lista interna, con los precandidatos de la agrupacion.
            formula = str(fila[0]).strip()
            resultados.append({**comun,
                               "agrupacion_id": agrupacion_actual["id"],
                               "agrupacion": agrupacion_actual["nombre"],
                               "formula": formula, "tipo_registro": "formula",
                               "votos": numero(fila[1]), "pct": numero(fila[2])})
            formulas.append({"anio": anio, "instancia": instancia,
                             "agrupacion_id": agrupacion_actual["id"],
                             "agrupacion": agrupacion_actual["nombre"],
                             "formula": formula, "votos": numero(fila[1])})

    return resultados, formulas


def verificar(resultados, archivo):
    """Las agrupaciones deben sumar el total de votos positivos."""
    suma = sum(f["votos"] for f in resultados
               if f["tipo_registro"] == "agrupacion"
               and isinstance(f["votos"], (int, float)))
    pos = next((f["votos"] for f in resultados
                if f["tipo_registro"] == "positivo"), None)
    if pos is None:
        return f"{archivo}: sin fila de VOTOS POSITIVOS"
    if abs(suma - pos) > 1:
        return (f"{archivo}: agrupaciones suman {suma:,.0f} != "
                f"positivos {pos:,.0f}")
    return None


def main():
    archivos = sorted(CRUDOS.glob("*.xlsx"))
    if not archivos:
        sys.exit(f"no hay .xlsx en {CRUDOS}")

    resultados, formulas, avisos = [], [], []
    for archivo in archivos:
        res, form = parsear(archivo)
        aviso = verificar(res, archivo.name)
        if aviso:
            avisos.append(aviso)
        resultados += res
        formulas += form
        print(f"  {archivo.name[:44]:44} {len(res):>3} filas, "
              f"{len(form)} fórmulas")

    SALIDA.mkdir(parents=True, exist_ok=True)
    for nombre, filas_datos, cols in (
        ("resultados_provincia_dine.csv", resultados, COLUMNAS),
        ("formulas.csv", formulas, COLUMNAS_FORMULAS),
    ):
        with open(SALIDA / nombre, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(filas_datos)
        print(f"-> datos/procesados/{nombre} ({len(filas_datos)} filas)")

    if avisos:
        print("\nControles de consistencia:")
        for a in avisos:
            print("  !", a)
    else:
        print("\nControles de consistencia: OK en todos los archivos.")


if __name__ == "__main__":
    main()
