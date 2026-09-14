"""Ingesta de las planillas DINE agregadas por distrito.

Estas planillas NO alimentan la tabla de hechos: estan agregadas a nivel
provincia y no traen departamento ni localidad (verificado hoja por hoja).
Sirven como *totales de control*: el total provincial por agrupacion que arroje
la base tiene que coincidir con ellos.

Cada libro trae una hoja por distrito mas una nacional. Se leen la hoja de
Santa Fe y la nacional. La instancia no se adivina del nombre del archivo: se
lee la fecha que la propia hoja declara en su encabezado y se la busca en el
calendario electoral.
"""

from __future__ import annotations

import re
import sys
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# Algunos libros de la DINE traen extensiones propietarias que openpyxl no
# conoce y avisa por cada hoja. El aviso no afecta los datos leidos.
warnings.filterwarnings("ignore", message="Unknown extension is not supported")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import comunes as c  # noqa: E402

_MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}

#: Etiquetas de cierre de cada planilla y su tipo de voto.
#: 'SUBTOTAL' marca las filas que son sumas de otras y no se cargan: si se
#: tomaran como una agrupacion mas, los votos se contarian dos veces. Le paso a
#: 2011-PASO, cuya hoja intercala 'VOTOS VALIDOS' (positivos + blancos).
_FOOTER = [
    ("VOTOS POSITIVOS", "POSITIVO"),
    ("VOTOS VALIDOS", "SUBTOTAL"),
    ("VOTOS EN BLANCO", "BLANCO"),
    ("VOTOS BLANCOS", "BLANCO"),
    ("VOTOS NULOS", "NULO"),
    ("VOTOS ANULADOS", "NULO"),
    ("VOTOS RECURRIDOS", "RECURRIDO"),
    ("VOTOS IMPUGNADOS", "IMPUGNADO"),
    ("VOTOS DE COMANDO", "IMPUGNADO"),
    ("VOTOS COMANDO", "IMPUGNADO"),
    ("TOTAL DE VOTANTES", "VOTANTES"),
    ("TOTAL DE VOTOS", "VOTANTES"),
]


@dataclass
class Control:
    fuente: dict
    totales: list[dict] = field(default_factory=list)
    padron: list[dict] = field(default_factory=list)


def _texto(valor) -> str:
    """Normaliza una celda a texto comparable: sin acentos, sin dobles espacios."""
    if valor is None:
        return ""
    # El espacio duro (\xa0) aparece en varias etiquetas de estas planillas.
    t = str(valor).replace("\xa0", " ")
    return re.sub(r"\s+", " ", c.sin_acentos(t)).strip().upper()


def _fecha_declarada(ws) -> str:
    """Lee del encabezado la fecha que la hoja declara, como 'AAAA-MM-DD'."""
    for i, fila in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 12:
            break
        for valor in fila:
            t = _texto(valor).lower()
            m = re.search(r"(\d{1,2})\s+de\s+([a-z]+)\s+de\s+(\d{4})", t)
            if m and m.group(2) in _MESES:
                return f"{int(m.group(3)):04d}-{_MESES[m.group(2)]:02d}-{int(m.group(1)):02d}"
    raise ValueError(f"La hoja {ws.title!r} no declara una fecha reconocible en su encabezado")


def _eleccion_por_fecha(fecha: str) -> str:
    for elec, (fecha_oficial, _orden, _nota) in c.CALENDARIO.items():
        if fecha_oficial == fecha:
            return elec
    raise ValueError(f"La fecha {fecha} no corresponde a ninguna instancia presidencial 2003-2023")


def _ubicar_encabezado(filas: list[tuple]) -> tuple[int, int, int, int | None]:
    """Encuentra la fila de encabezado y las columnas de etiqueta, votos y formula.

    Devuelve (indice_fila, col_etiqueta, col_votos, col_formula_o_None).
    El layout cambia entre anios, asi que nunca se asumen posiciones fijas.
    """
    for i, fila in enumerate(filas):
        col_votos = col_etiqueta = col_formula = None
        for j, valor in enumerate(fila):
            t = _texto(valor)
            if t == "VOTOS":
                col_votos = j
            elif "AGRUPACION" in t:
                col_etiqueta = j
            elif "FORMULA" in t:
                # '_FORMULAS_' sola (2003) o 'FORMULAS - AGRUPACIONES' (2019).
                if "AGRUPACION" in t:
                    col_etiqueta = j
                else:
                    col_formula = j
        if col_votos is not None and col_etiqueta is not None:
            return i, col_etiqueta, col_votos, col_formula
    raise ValueError("No se encontro la fila de encabezado (VOTOS + AGRUPACIONES)")


def _leer_hoja(ws, ambito: str, distrito: str, fuente_id: str) -> Control:
    filas = list(ws.iter_rows(values_only=True))
    elec = _eleccion_por_fecha(_fecha_declarada(ws))
    i_enc, col_etq, col_votos, col_formula = _ubicar_encabezado(filas)
    cuerpo = filas[i_enc + 1:]

    # En el layout 2019 cada agrupacion lleva un numero en la primera columna y
    # debajo va una fila de lista/candidatos que no hay que sumar. Si ninguna
    # fila trae ese numero (layout 2003), se toman todas.
    def _id_agrupacion(fila):
        valor = fila[0] if fila else None
        t = _texto(valor)
        return t if re.fullmatch(r"\d+", t) else None

    hay_ids = any(_id_agrupacion(f) for f in cuerpo[:25])

    control = Control(fuente={
        "fuente_id": fuente_id,
        "eleccion_id": elec,
        "ambito": ambito,
        "distrito": distrito,
    })

    for fila in cuerpo:
        etiqueta = _texto(fila[col_etq]) if col_etq < len(fila) else ""
        if not etiqueta:
            continue

        cierre = next((tipo for pref, tipo in _FOOTER if etiqueta.startswith(pref)), None)
        if cierre:
            cantidad = c.a_entero(fila[col_votos]) if col_votos < len(fila) else 0
            if cierre == "SUBTOTAL":
                continue
            if cierre == "VOTANTES":
                control.padron.append({
                    "eleccion_id": elec, "ambito": ambito, "distrito": distrito,
                    "votantes": cantidad, "fuente_id": fuente_id,
                })
            elif cierre != "POSITIVO":
                # El positivo total se recalcula sumando agrupaciones; cargarlo
                # aca lo duplicaria. Queda como control en el validador.
                control.totales.append({
                    "eleccion_id": elec, "ambito": ambito, "distrito": distrito,
                    "tipo_voto": cierre, "agrupacion_key": None,
                    "nombre_fuente": None, "formula": None,
                    "votos": cantidad, "fuente_id": fuente_id,
                })
            continue

        # Guardarrail: una etiqueta de total que no este en el vocabulario no
        # puede entrar como si fuera un partido. Antes que inventar una
        # agrupacion, el ingestor falla y nombra la etiqueta desconocida.
        if etiqueta.startswith("VOTOS ") or etiqueta.startswith("TOTAL "):
            raise ValueError(
                f"Etiqueta de total no reconocida: {etiqueta!r}. "
                "Agregarla a _FOOTER antes de seguir; si se la deja pasar se "
                "carga como una agrupacion inexistente.")

        if hay_ids and not _id_agrupacion(fila):
            continue  # fila de lista o de candidatos: detalle, no total
        if col_votos >= len(fila) or fila[col_votos] is None:
            continue

        nombre = c.normalizar_nombre(str(fila[col_etq]).replace("\xa0", " "))
        formula = None
        if col_formula is not None and col_formula < len(fila) and fila[col_formula]:
            formula = c.normalizar_nombre(str(fila[col_formula]))
        control.totales.append({
            "eleccion_id": elec, "ambito": ambito, "distrito": distrito,
            "tipo_voto": "POSITIVO", "agrupacion_key": c.clave(nombre),
            "nombre_fuente": nombre, "formula": formula,
            "votos": c.a_entero(fila[col_votos]), "fuente_id": fuente_id,
        })

    # Electores y mesas del encabezado.
    electores = mesas = None
    for fila in filas[:i_enc]:
        etiqueta = _texto(fila[0]) or (_texto(fila[1]) if len(fila) > 1 else "")
        numero = next((v for v in fila[1:] if isinstance(v, (int, float))), None)
        if numero is None:
            continue
        if "ELECTORES" in etiqueta:
            electores = c.a_entero(numero)
        elif etiqueta.startswith("MESAS"):
            mesas = c.a_entero(numero)
    for fila_padron in control.padron:
        fila_padron["electores"] = electores
        fila_padron["mesas"] = mesas

    _verificar_cuadre(control)
    return control


def _verificar_cuadre(control: Control) -> None:
    """La suma de positivos y no positivos tiene que dar el total de votantes.

    Es el control que detecta que una fila de subtotal se haya colado como
    agrupacion, o que falte un tipo de voto. Falla en vez de avisar: una
    planilla de control que no cuadra no sirve para controlar nada.
    """
    if not control.padron:
        return
    votantes = control.padron[0].get("votantes")
    if not votantes:
        return
    suma = sum(t["votos"] for t in control.totales)
    if suma != votantes:
        raise ValueError(
            f"{control.fuente['eleccion_id']} / {control.fuente['distrito']}: "
            f"los votos suman {suma:,} y la planilla declara {votantes:,} votantes "
            f"(diferencia {suma - votantes:+,})")


def leer_libro(ruta: Path) -> tuple[list[Control], list[dict]]:
    """Lee las hojas de Santa Fe y nacional de una planilla DINE por distrito.

    Devuelve los controles leidos y la lista de hojas omitidas con su motivo.
    Algunos libros traen una hoja nacional que es un indice de formulas y
    partidos componentes, sin columna de votos: se omite, pero queda anotada.
    Si la que falla es la hoja de Santa Fe, es un error: es la que importa.
    """
    ruta = Path(ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True)
    sha = c.sha256_archivo(ruta)
    salida: list[Control] = []
    omitidas: list[dict] = []
    for nombre in wb.sheetnames:
        t = _texto(nombre)
        if t in ("SANTA FE",):
            ambito, distrito = "PROVINCIA", "Santa Fe"
        elif t.startswith("NACIONAL") or t.startswith("TOTAL PAIS"):
            ambito, distrito = "PAIS", "Total pais"
        else:
            continue
        ws = wb[nombre]
        try:
            control = _leer_hoja(ws, ambito, distrito, f"CONTROL_{ruta.stem}_{c.clave(nombre)}")
        except ValueError as err:
            if ambito == "PROVINCIA":
                raise ValueError(f"{ruta.name} / hoja {nombre!r}: {err}") from err
            omitidas.append({"archivo": ruta.name, "hoja": nombre, "motivo": str(err)})
            continue
        control.fuente.update({
            "archivo": ruta.name,
            "organismo": "Dirección Nacional Electoral (DINE)",
            "unidad_original": "Distrito (provincia). No trae departamento ni localidad.",
            "sha256": sha,
        })
        salida.append(control)
    wb.close()
    return salida, omitidas
