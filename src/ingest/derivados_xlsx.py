"""Ingesta de los archivos 'Votos por Localidad - <instancia> - Presidente - Santa Fe.xlsx'.

Son productos ya procesados en etapas previas del proyecto, a nivel localidad.
Cada libro trae seis hojas; se usan cinco:

  Votos por Localidad            -> votos POSITIVO por agrupacion/formula
  Totales por Tipo de Voto       -> BLANCO / NULO / RECURRIDO / IMPUGNADO
                                    (su columna POSITIVO se usa para validar)
  Electores y Mesas              -> padron
  Nomenclador Circuito-Localidad -> dim_circuito
  Metodologia                    -> registro de fuente

La hoja 'Agrupacion Ganadora por Localidad' se ignora: es derivable.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import comunes as c  # noqa: E402

#: Como se nombra cada instancia en el nombre de archivo.
_INSTANCIA_EN_NOMBRE = [
    (r"\bPASO\b", "PASO"),
    (r"\bGenerales?\b", "GENERAL"),
    (r"\bBallotage\b", "BALLOTAGE"),
    (r"\bSegunda\s+Vuelta\b", "BALLOTAGE"),
]


@dataclass
class Ingesta:
    """Resultado de leer un libro: filas listas para consolidar."""

    fuente: dict
    hechos: list[dict] = field(default_factory=list)
    padron: list[dict] = field(default_factory=list)
    circuitos: list[dict] = field(default_factory=list)
    incidencias: list[dict] = field(default_factory=list)


def _hoja(wb, prefijo: str):
    """Busca una hoja por prefijo, sin acentos ni mayusculas.

    Excel trunca los nombres de hoja a 31 caracteres, asi que nunca se compara
    por igualdad exacta.
    """
    objetivo = c.sin_acentos(prefijo).upper()
    for ws in wb.worksheets:
        if c.sin_acentos(ws.title).upper().startswith(objetivo):
            return ws
    raise KeyError(f"No se encontro la hoja que empieza con {prefijo!r}; hay {wb.sheetnames}")


def _filas(ws) -> list[dict]:
    """Lee una hoja con encabezado en la primera fila y devuelve dicts.

    Las claves se normalizan (mayusculas, sin acentos) para no depender de como
    escribio el encabezado cada archivo.
    """
    it = ws.iter_rows(values_only=True)
    try:
        crudo = next(it)
    except StopIteration:
        return []
    encabezado = [c.sin_acentos(str(h)).upper().strip() if h is not None else "" for h in crudo]
    salida = []
    for fila in it:
        if all(v is None or str(v).strip() == "" for v in fila):
            continue
        salida.append({k: v for k, v in zip(encabezado, fila) if k})
    return salida


def identificar_eleccion(nombre_archivo: str) -> str:
    """Deduce el eleccion_id del nombre del archivo.

    'Votos por Localidad - Ballotage 2015 - ...' -> '2015-BALLOTAGE'
    'Votos por Localidad - 2007 - ...'           -> '2007-GENERAL'
    """
    anios = re.findall(r"\b(20\d{2})\b", nombre_archivo)
    if not anios:
        raise ValueError(f"No se pudo leer el anio en {nombre_archivo!r}")
    anio = int(anios[0])
    instancia = "GENERAL"
    for patron, valor in _INSTANCIA_EN_NOMBRE:
        if re.search(patron, nombre_archivo, flags=re.IGNORECASE):
            instancia = valor
            break
    return c.eleccion_id(anio, instancia)


def _geo(fila: dict) -> tuple[str, str, str, str]:
    """Extrae (departamento_id, departamento, localidad_id, localidad) de una fila."""
    depto = c.normalizar_nombre(fila.get("DEPARTAMENTO"))
    loc = c.normalizar_nombre(fila.get("LOCALIDAD"))
    return c.clave(depto), depto, c.clave(loc), loc


def leer_libro(ruta: Path) -> Ingesta:
    """Lee un libro completo y devuelve sus filas normalizadas."""
    ruta = Path(ruta)
    elec = identificar_eleccion(ruta.name)
    etiqueta = c.ETIQUETA_POR_ELECCION[elec]
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    metodologia = {
        c.normalizar_nombre(f.get("CONCEPTO")): f.get("DESCRIPCION")
        for f in _filas(_hoja(wb, "Metodologia"))
        if f.get("CONCEPTO")
    }
    fuente_id = f"DERIVADO_{elec}"
    fuente = {
        "fuente_id": fuente_id,
        "archivo": ruta.name,
        "eleccion_id": elec,
        "organismo": metodologia.get("Organismo de origen"),
        "archivo_origen": metodologia.get("Archivo fuente requerido"),
        "unidad_original": metodologia.get("Unidad original del archivo fuente"),
        "recuento": "NO_DECLARADO",
        "cobertura_declarada": metodologia.get("Departamentos cubiertos")
        or metodologia.get("Departamentos cubiertos en este archivo"),
        "fecha_proceso_origen": metodologia.get("Fecha de procesamiento"),
        "sha256": c.sha256_archivo(ruta),
        "observaciones": metodologia.get("Particularidad de esta instancia electoral")
        or metodologia.get("Particularidad respecto de archivos DINE 2011 en adelante"),
    }

    ing = Ingesta(fuente=fuente)

    # --- Hoja 1: votos positivos por agrupacion/formula --------------------
    positivos_por_localidad: dict[str, int] = {}
    for fila in _filas(_hoja(wb, "Votos por Localidad")):
        depto_id, depto, loc_id, loc = _geo(fila)
        etiqueta_fuente = c.normalizar_nombre(fila.get("AGRUPACION_POLITICA"))
        votos = c.a_entero(fila.get("VOTOS"))
        ing.hechos.append({
            "eleccion_id": elec,
            "departamento_id": depto_id,
            "departamento": depto,
            "localidad_id": loc_id,
            "localidad": loc,
            "tipo_voto": "POSITIVO",
            "agrupacion_key": c.clave(etiqueta_fuente),
            "agrupacion_nombre_fuente": etiqueta_fuente,
            "formula": etiqueta_fuente if etiqueta == "FORMULA" else None,
            "votos": votos,
            "fuente_id": fuente_id,
        })
        positivos_por_localidad[loc_id] = positivos_por_localidad.get(loc_id, 0) + votos

    # --- Hoja 3: tipos de voto -------------------------------------------
    # Solo se cargan los NO positivos: el positivo ya vino desagregado arriba y
    # cargarlo de nuevo lo duplicaria. La columna POSITIVO se usa para validar.
    ws_tipos = _hoja(wb, "Totales por Tipo de Voto")
    for fila in _filas(ws_tipos):
        depto_id, depto, loc_id, loc = _geo(fila)
        for columna, valor in fila.items():
            if columna in ("DEPARTAMENTO", "LOCALIDAD"):
                continue
            tipo = c.normalizar_tipo_voto(columna)
            cantidad = c.a_entero(valor)
            if tipo == "POSITIVO":
                declarado = cantidad
                sumado = positivos_por_localidad.get(loc_id, 0)
                if declarado != sumado:
                    ing.incidencias.append({
                        "eleccion_id": elec,
                        "localidad_id": loc_id,
                        "control": "positivos_declarados_vs_suma_agrupaciones",
                        "esperado": declarado,
                        "obtenido": sumado,
                        "diferencia": sumado - declarado,
                    })
                continue
            ing.hechos.append({
                "eleccion_id": elec,
                "departamento_id": depto_id,
                "departamento": depto,
                "localidad_id": loc_id,
                "localidad": loc,
                "tipo_voto": tipo,
                "agrupacion_key": None,
                "agrupacion_nombre_fuente": None,
                "formula": None,
                "votos": cantidad,
                "fuente_id": fuente_id,
            })

    # --- Hoja 4: padron ---------------------------------------------------
    for fila in _filas(_hoja(wb, "Electores y Mesas")):
        depto_id, depto, loc_id, loc = _geo(fila)
        ing.padron.append({
            "eleccion_id": elec,
            "departamento_id": depto_id,
            "localidad_id": loc_id,
            "mesas": c.a_entero(fila.get("MESAS")),
            "electores": c.a_entero(fila.get("ELECTORES")),
            "fuente_id": fuente_id,
        })

    # --- Hoja 5: nomenclador ---------------------------------------------
    for fila in _filas(_hoja(wb, "Nomenclador Circuito")):
        depto_id, depto, loc_id, loc = _geo(fila)
        circuito = c.normalizar_nombre(fila.get("CIRCUITO"))
        ing.circuitos.append({
            "circuito_id": circuito,
            "departamento_id": depto_id,
            "departamento": depto,
            "localidad_id": loc_id,
            "localidad": loc,
            "eleccion_id": elec,
            "fuente_id": fuente_id,
        })

    wb.close()
    return ing
